"""Read-only LPM preprocessing and cohort-relative RPI analysis.

Private source paths come from a local manifest. Outputs are restricted to
workspace/private. Names are grouping strings, not verified identities.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
from io import BytesIO
import hashlib
import importlib.metadata
import json
import math
from pathlib import Path
import re
import subprocess
import sys
import tomllib

import numpy as np
from openpyxl import load_workbook
from scipy import stats

DIMENSIONS = {
    "Pedagogical": tuple(range(0, 6)),
    "Professional": tuple(range(6, 11)),
    "Personality": tuple(range(11, 16)),
    "Social": tuple(range(16, 20)),
}
ROOT = Path(__file__).resolve().parents[1]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def parse_item(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("missing_or_invalid_item")
    try:
        number = float(value)
    except (ValueError, TypeError):
        raise ValueError("nonnumeric_item") from None
    if not math.isfinite(number) or not 1 <= number <= 5:
        raise ValueError("item_outside_verified_scale")
    return number


def exclusion_reason(name, config):
    label = " ".join(name.split()).upper()
    rules = config["exclusions"]
    if re.match(rules["team_regex"], label):
        return "team_label"
    if label in rules["generic_supervisor_labels"]:
        return "generic_supervisor_label"
    if label in rules["other_generic_labels"]:
        return "other_generic_label"
    return None


def read_sources(manifest):
    records, sources = [], []
    required = [f"Pertanyaan {i}" for i in range(1, 21)]
    for entry in manifest:
        path = Path(entry["path"])
        payload = path.read_bytes()
        digest = hashlib.sha256(payload).hexdigest()
        if digest != entry["sha256"]:
            raise ValueError("Source hash differs from reviewed manifest")
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        tables = 0
        count = 0
        for sheet in workbook:
            columns = None
            for row_number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                values = list(cells)
                if "Nama Dosen" in values:
                    if all(item in values for item in required):
                        columns = {name: values.index(name) for name in required + ["Nama Dosen", "Program Studi", "No."]}
                        tables += 1
                    else:
                        columns = None
                    continue
                if columns is None or not any(v is not None for v in values):
                    continue
                number = values[columns["No."]]
                name = values[columns["Nama Dosen"]]
                program = values[columns["Program Studi"]]
                if not isinstance(number, (int, float)) or isinstance(number, bool):
                    raise ValueError("Unexpected nonrecord row inside item table")
                if not isinstance(name, str) or not name.strip() or not isinstance(program, str) or not program.strip():
                    raise ValueError("Missing lecturer/program label")
                records.append({"name": name, "program": program,
                                "items": [values[columns[item]] for item in required],
                                "source": {"file": path.name, "sha256": digest,
                                           "sheet": sheet.title, "row": row_number}})
                count += 1
        workbook.close()
        if tables != 1:
            raise ValueError("Expected exactly one item table per reviewed workbook")
        if sha256(path) != digest:
            raise ValueError("Source changed during read")
        sources.append({"path": str(path), "sha256": digest, "records": count})
    return records, sources


def prepare(records, config):
    pairs = Counter((r["name"], r["program"]) for r in records)
    if any(count > 1 for count in pairs.values()):
        raise ValueError("Duplicate name-program pairs require source review")
    retained, excluded = [], []
    for record in records:
        if len(record["items"]) != 20:
            raise ValueError("Expected exactly 20 items")
        reasons = []
        reason = exclusion_reason(record["name"], config)
        if reason:
            reasons.append(reason)
        values = []
        invalid = []
        for index, value in enumerate(record["items"], 1):
            try:
                values.append(parse_item(value))
            except ValueError as error:
                invalid.append({"item": index, "reason": str(error)})
        if invalid:
            reasons.append("incomplete_or_invalid_record")
        if reasons:
            excluded.append(dict(record, reasons=reasons, invalid_items=invalid))
        else:
            retained.append(dict(record, items=values, total=sum(values)))
    assert len(records) == len(retained) + len(excluded)
    return retained, excluded


def aggregate(records):
    groups = defaultdict(list)
    for record in records:
        groups[record["name"]].append(record)
    if not groups:
        raise ValueError("No eligible records")
    means, identities = [], []
    for index, name in enumerate(sorted(groups), 1):
        group = groups[name]
        means.append(np.mean([r["items"] for r in group], axis=0))
        identities.append({"code": f"D{index:03d}", "original_name": name,
                           "n_records": len(group),
                           "source_records": [r["source"] for r in group]})
    return np.array(means), identities


def validate_weights(weights, atol=1e-12):
    weights = np.asarray(weights, dtype=float)
    if weights.shape != (4,) or not np.isfinite(weights).all() or (weights < 0).any():
        raise ValueError("Four finite nonnegative weights are required")
    if not math.isclose(float(weights.sum()), 1.0, rel_tol=0, abs_tol=atol):
        raise ValueError("Weights must sum to one")
    return weights


def score(item_means, config):
    matrix = np.asarray(item_means, dtype=float)
    if matrix.ndim != 2 or matrix.shape[1] != 20 or len(matrix) == 0:
        raise ValueError("Expected nonempty entity-by-20 item matrix")
    if not np.isfinite(matrix).all() or (matrix < 1).any() or (matrix > 5).any():
        raise ValueError("Invalid aggregated item score")
    expected = {"rank_ties": "exact_unrounded", "percentiles": "100*(ascending_midrank-0.5)/N",
                "quantiles": "linear", "quadrants": "median_equal_is_high",
                "skewness_bias": False, "kurtosis_fisher": True, "kurtosis_bias": False}
    if any(config.get(key) != value for key, value in expected.items()):
        raise ValueError("Unsupported analysis convention")
    weights = validate_weights(config["weights"], config["weight_sum_atol"])
    offset, location, scale = (float(config[k]) for k in ["blom_offset", "location", "scale"])
    if not all(math.isfinite(x) for x in [offset, location, scale]) or not 0 <= offset < 1 or scale <= 0:
        raise ValueError("Invalid plotting-position/display constants")
    n = len(matrix)
    dimensions = np.column_stack([matrix[:, indices].mean(axis=1) for indices in DIMENSIONS.values()])
    ranks = np.column_stack([stats.rankdata(dimensions[:, k], method="average") for k in range(4)])
    p = (ranks - offset) / (n + 1 - 2 * offset)
    if not ((p > 0) & (p < 1)).all():
        raise ValueError("Plotting positions must be strictly inside (0,1)")
    z = stats.norm.ppf(p)
    composite = z @ weights
    rpi = location + scale * composite
    item_sum = matrix.sum(axis=1)
    matched_raw = 20 * (dimensions @ weights)
    rank_raw = stats.rankdata(-matched_raw, method="min")
    rank_rpi = stats.rankdata(-rpi, method="min")
    rank_item_sum = stats.rankdata(-item_sum, method="min")
    percentile = lambda values: 100 * (stats.rankdata(values, method="average") - 0.5) / n
    med_raw, med_rpi = np.median(matched_raw), np.median(rpi)
    quadrants = ["Q1" if x >= med_raw and y >= med_rpi else
                 "Q2" if x >= med_raw and y < med_rpi else
                 "Q3" if x < med_raw and y >= med_rpi else "Q4"
                 for x, y in zip(matched_raw, rpi)]
    return {"item_means": matrix, "dimensions": dimensions, "dimension_ranks": ranks,
            "p": p, "z": z, "C": composite, "RPI": rpi,
            "item_sum": item_sum, "matched_raw": matched_raw,
            "rank_item_sum": rank_item_sum, "rank_matched_raw": rank_raw,
            "rank_RPI": rank_rpi, "matched_rank_shift": rank_raw - rank_rpi,
            "item_sum_rank_shift": rank_item_sum - rank_rpi,
            "percentile_matched_raw": percentile(matched_raw), "percentile_RPI": percentile(rpi),
            "quadrants": quadrants, "weights": weights}


def finite_or_none(value):
    return float(value) if math.isfinite(float(value)) else None


def describe(values, maximum=None):
    x = np.asarray(values, dtype=float)
    n = len(x)
    unique, counts = np.unique(x, return_counts=True)
    nonconstant = len(unique) > 1
    q1, q3 = np.quantile(x, [0.25, 0.75], method="linear")
    shapiro = stats.shapiro(x) if nonconstant and 3 <= n <= 5000 else None
    return {"N": n, "mean": float(np.mean(x)), "median": float(np.median(x)),
            "SD": float(np.std(x, ddof=1)) if n > 1 else None,
            "min": float(x.min()), "max": float(x.max()), "IQR": float(q3-q1),
            "skewness": finite_or_none(stats.skew(x, bias=False)) if nonconstant and n >= 3 else None,
            "excess_kurtosis": finite_or_none(stats.kurtosis(x, fisher=True, bias=False)) if nonconstant and n >= 4 else None,
            "unique_values": len(unique), "records_in_ties": int(counts[counts > 1].sum()),
            "largest_tie": int(counts.max()), "theoretical_max": maximum,
            "at_theoretical_max": int(np.sum(x == maximum)) if maximum is not None else None,
            "shapiro_W": float(shapiro.statistic) if shapiro else None,
            "shapiro_p": float(shapiro.pvalue) if shapiro else None,
            "shapiro_status": "computed" if shapiro else "unavailable_constant_or_N_outside_3_to_5000"}


def cronbach_alpha(matrix):
    x = np.asarray(matrix, dtype=float)
    if x.ndim != 2 or x.shape[1] < 2 or len(x) < 2 or not np.isfinite(x).all():
        return None
    variance = np.var(x.sum(axis=1), ddof=1)
    if variance <= 0:
        return None
    k = x.shape[1]
    return float(k/(k-1) * (1 - np.var(x, axis=0, ddof=1).sum()/variance))


def agreement(x, y, shift):
    valid = len(x) >= 2 and np.ptp(x) > 0 and np.ptp(y) > 0
    rho = stats.spearmanr(x, y) if valid else None
    tau = stats.kendalltau(x, y, variant="b") if valid else None
    return {"Spearman_rho": float(rho.statistic) if rho else None,
            "Spearman_p": float(rho.pvalue) if rho else None,
            "Kendall_tau_b": float(tau.statistic) if tau else None,
            "Kendall_p": float(tau.pvalue) if tau else None,
            "mean_abs_rank_shift": float(np.mean(np.abs(shift))),
            "max_abs_rank_shift": int(np.max(np.abs(shift))),
            "n_position_changed": int(np.count_nonzero(shift))}


def diagnostics(scored, retained, config):
    series = {f"P{i+1}": describe(scored["item_means"][:, i], 5) for i in range(20)}
    for k, name in enumerate(DIMENSIONS):
        series[name] = describe(scored["dimensions"][:, k], 5)
    for name in ["item_sum", "matched_raw", "RPI"]:
        series[name] = describe(scored[name], 100 if name != "RPI" else None)
    reliability = {}
    matrices = {"lecturer_program_aggregate_records": np.array([r["items"] for r in retained]),
                "lecturer_name_group_item_means": scored["item_means"]}
    for level, matrix in matrices.items():
        reliability[level] = {"N": len(matrix), "missing_policy": "complete_record_only",
                              "overall_20_items": cronbach_alpha(matrix),
                              "dimensions": {name: cronbach_alpha(matrix[:, indices])
                                             for name, indices in DIMENSIONS.items()}}
    dimensions = scored["dimensions"]
    pearson = [[float(stats.pearsonr(dimensions[:, i], dimensions[:, j]).statistic)
                if len(dimensions) >= 2 and np.ptp(dimensions[:, i]) and np.ptp(dimensions[:, j]) else None
                for j in range(4)] for i in range(4)]
    n = len(dimensions)
    predicted_sd = None
    observed_sd = float(np.std(scored["RPI"], ddof=1)) if n > 1 else None
    if n > 1:
        covariance = np.cov(scored["z"], rowvar=False, ddof=1)
        variance = float(scored["weights"] @ covariance @ scored["weights"])
        predicted_sd = config["scale"] * math.sqrt(max(0.0, variance))
        if not math.isclose(predicted_sd, observed_sd, rel_tol=0, abs_tol=config["covariance_check_atol"]):
            raise ValueError("Composite covariance identity failed")
    return {"analysis_status": "provisional_source_name_groups", "series": series,
            "reliability": reliability, "pearson_dimension_order": list(DIMENSIONS),
            "pearson_dimensions": pearson,
            "matched_weight_agreement": agreement(scored["matched_raw"], scored["RPI"], scored["matched_rank_shift"]),
            "item_sum_agreement": agreement(scored["item_sum"], scored["RPI"], scored["item_sum_rank_shift"]),
            "RPI_SD_observed": observed_sd, "RPI_SD_from_covariance": predicted_sd,
            "quadrants": dict(Counter(scored["quadrants"])),
            "quadrant_definitions": {"Q1": "High raw and high RPI", "Q2": "High raw and low RPI",
                                      "Q3": "Low raw and high RPI", "Q4": "Low raw and low RPI"},
            "upper_range_counts": {name: {str(t): int(np.sum(scored[name] >= t)) for t in [85,90,95]}
                                   for name in ["item_sum", "matched_raw"]}}


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def run(manifest_path, cohort_path, analysis_path, output):
    output = Path(output).resolve()
    private = (ROOT / "private").resolve()
    if not output.is_relative_to(private):
        raise ValueError("Institutional outputs must be inside workspace/private")
    if output.exists() and any(output.iterdir()):
        raise ValueError("Output directory must be new or empty; preserve prior runs")
    cohort = tomllib.loads(Path(cohort_path).read_text(encoding="utf-8"))
    analysis = tomllib.loads(Path(analysis_path).read_text(encoding="utf-8"))
    if cohort["status"] != "provisional" or cohort["entity_key"] != "original_source_name_string" or cohort["record_weighting"] != "equal":
        raise ValueError("Only the documented provisional equal-record scenario is implemented")
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    records, sources = read_sources(manifest)
    retained, excluded = prepare(records, cohort)
    item_means, identities = aggregate(retained)
    scored = score(item_means, analysis)
    report = diagnostics(scored, retained, analysis)
    score_rows = []
    for i, identity in enumerate(identities):
        row = {"code": identity["code"], "n_records": identity["n_records"]}
        for key, value in scored.items():
            if key == "weights":
                continue
            if key == "quadrants":
                row[key] = value[i]
            else:
                entry = value[i]
                row[key] = entry.tolist() if isinstance(entry, np.ndarray) else float(entry)
        score_rows.append(row)
    output.mkdir(parents=True, exist_ok=True)
    outputs = {"retained_records.json": retained, "exclusions.json": excluded,
               "identity_mapping.json": identities, "scores.json": score_rows,
               "diagnostics.json": report}
    for name, value in outputs.items():
        write_json(output / name, value)
    run_manifest = {"timestamp": datetime.now(timezone(timedelta(hours=7))).isoformat(),
                    "timezone": "Asia/Bangkok", "status": "provisional; not roster verified",
                    "sources": sources, "config": {"cohort": cohort, "analysis": analysis},
                    "config_hashes": {str(p): sha256(p) for p in [cohort_path, analysis_path]},
                    "code_hashes": {str(p.relative_to(ROOT)): sha256(p) for p in sorted((ROOT/"src").glob("*.py"))},
                    "commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
                    "working_tree_dirty": bool(subprocess.check_output(["git", "status", "--porcelain"], cwd=ROOT, text=True).strip()),
                    "python": sys.version,
                    "dependencies": {name: importlib.metadata.version(name) for name in ["numpy", "scipy", "openpyxl", "et_xmlfile"]},
                    "counts": {"input_records": len(records), "retained_records": len(retained),
                               "excluded_records": len(excluded), "name_groups": len(identities),
                               "exclusion_reasons": dict(Counter(reason for r in excluded for reason in r["reasons"]))},
                    "output_hashes": {name: sha256(output/name) for name in outputs}}
    write_json(output/"manifest.json", run_manifest)
    return run_manifest["counts"], report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-manifest", required=True)
    parser.add_argument("--cohort", default="configs/cohort.toml")
    parser.add_argument("--analysis", default="configs/analysis.toml")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    counts, report = run(args.source_manifest, args.cohort, args.analysis, args.output)
    print(json.dumps({"counts": counts, "matched_weight_agreement": report["matched_weight_agreement"],
                      "RPI_SD": report["RPI_SD_observed"], "status": report["analysis_status"]}, indent=2))


if __name__ == "__main__":
    main()
